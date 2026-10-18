"""Isolated attachment parser (§05 step 5 "mengisolasi parser", §25 "parser isolation").

Runs in a separate interpreter started by ``aeos.intake.parser`` as::

    python -I parse_worker.py <input> <output.json> <kind> <max_chars>

(``python -I -m aeos.intake.parse_worker`` works too). ``-I`` ignores PYTHON* environment
variables, the user site directory and the script directory, so nothing from the upload location
or the caller's environment is imported. This module deliberately imports nothing from ``aeos``.
It only *reads* the input file and writes one JSON document; it never executes, renders or
follows links in the content. The parent enforces timeout, memory and output-size caps and kills
the process tree on violation.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

MAX_PDF_PAGES = 500
# Self-imposed POSIX limits (the parent also enforces wall-clock/RSS/output caps on every OS).
LIMIT_AS_BYTES = 2 * 1024 * 1024 * 1024
LIMIT_CPU_SECONDS = 35
LIMIT_FSIZE_BYTES = 16 * 1024 * 1024


def _apply_limits() -> None:
    if sys.platform == "win32":
        return
    try:
        import resource
    except ImportError:  # pragma: no cover
        return
    for name, value in (("RLIMIT_AS", LIMIT_AS_BYTES), ("RLIMIT_CPU", LIMIT_CPU_SECONDS),
                        ("RLIMIT_FSIZE", LIMIT_FSIZE_BYTES)):
        res = getattr(resource, name, None)
        if res is None:
            continue
        try:
            _soft, hard = resource.getrlimit(res)
            cap = value if hard == resource.RLIM_INFINITY else min(value, hard)
            resource.setrlimit(res, (cap, hard))
        except (ValueError, OSError):  # e.g. RLIMIT_AS unsupported on macOS
            continue


def _text_file(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    return data.decode("utf-8-sig", errors="replace")


def _parse_text(path: Path, kind: str) -> tuple[str, dict[str, Any]]:
    text = _text_file(path.read_bytes())
    meta: dict[str, Any] = {"lines": text.count("\n") + (1 if text else 0)}
    if kind == "json":
        try:
            json.loads(text)
            meta["valid_json"] = True
        except (ValueError, RecursionError) as exc:
            meta["valid_json"] = False
            meta["json_error"] = str(exc)[:200]
    elif kind == "yaml":
        try:
            import yaml

            list(yaml.safe_load_all(text))
            meta["valid_yaml"] = True
        except Exception as exc:
            meta["valid_yaml"] = False
            meta["yaml_error"] = str(exc).splitlines()[0][:200] if str(exc) else exc.__class__.__name__
    elif kind == "csv":
        import csv
        import io

        try:
            rows = list(csv.reader(io.StringIO(text)))
            meta["rows"] = len(rows)
            meta["columns"] = max((len(r) for r in rows), default=0)
        except csv.Error as exc:
            meta["csv_error"] = str(exc)[:200]
    return text, meta


def _parse_docx(path: Path) -> tuple[str, dict[str, Any]]:
    import docx  # python-docx; lxml parser with entity resolution disabled

    doc = docx.Document(str(path))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text]
    tables = 0
    for table in doc.tables:
        tables += 1
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts), {"paragraphs": len(doc.paragraphs), "tables": tables}


def _parse_pdf(path: Path) -> tuple[str, dict[str, Any]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ValueError("encrypted PDF cannot be parsed")
    pages = reader.pages
    out: list[str] = []
    total = len(pages)
    for i, page in enumerate(pages):
        if i >= MAX_PDF_PAGES:
            break
        try:
            out.append(page.extract_text() or "")
        except Exception as exc:  # one broken page must not hide the rest
            out.append(f"[page {i + 1} unreadable: {exc.__class__.__name__}]")
    return "\n\n".join(out), {"pages": total, "pages_parsed": min(total, MAX_PDF_PAGES)}


def _parse_zip(path: Path) -> tuple[str, dict[str, Any]]:
    import zipfile

    with zipfile.ZipFile(str(path)) as zf:  # central directory only; nothing is extracted
        infos = zf.infolist()
        lines = [f"- {i.filename} ({i.file_size} bytes)" for i in infos]
    return "Archive members (not extracted):\n" + "\n".join(lines), {"members": len(infos)}


def _image_meta(path: Path) -> dict[str, Any]:
    head = path.read_bytes()[:64]
    meta: dict[str, Any] = {"bytes": path.stat().st_size}
    if head.startswith(b"\x89PNG\r\n\x1a\n") and len(head) >= 24:
        meta["width"] = int.from_bytes(head[16:20], "big")
        meta["height"] = int.from_bytes(head[20:24], "big")
    elif head[:6] in (b"GIF87a", b"GIF89a") and len(head) >= 10:
        meta["width"] = int.from_bytes(head[6:8], "little")
        meta["height"] = int.from_bytes(head[8:10], "little")
    return meta


def parse(path: Path, kind: str, max_chars: int) -> dict[str, Any]:
    if kind == "image":
        return {"status": "SKIPPED", "text": "", "meta": _image_meta(path),
                "error": "images are stored with metadata only (no OCR)"}
    if kind in {"markdown", "text", "json", "yaml", "csv"}:
        text, meta = _parse_text(path, kind)
    elif kind == "docx":
        text, meta = _parse_docx(path)
    elif kind == "pdf":
        text, meta = _parse_pdf(path)
    elif kind == "zip":
        text, meta = _parse_zip(path)
    else:
        return {"status": "FAILED", "text": "", "meta": {}, "error": f"unsupported kind {kind}"}
    meta["chars"] = len(text)
    meta["truncated"] = len(text) > max_chars
    return {"status": "PARSED", "text": text[:max_chars], "meta": meta, "error": ""}


def main(argv: list[str], *, limits: bool = True) -> int:
    if len(argv) != 4:
        sys.stderr.write("usage: parse_worker.py <input> <output.json> <kind> <max_chars>\n")
        return 2
    src, dst, kind, max_chars = Path(argv[0]), Path(argv[1]), argv[2], int(argv[3])
    if limits:  # only in the dedicated worker process, never when imported in-process
        _apply_limits()
    try:
        result = parse(src, kind, max_chars)
    except Exception as exc:
        result = {"status": "FAILED", "text": "", "meta": {}, "error": f"{exc.__class__.__name__}: {str(exc)[:300]}"}
    tmp = dst.with_suffix(".partial")
    tmp.write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
    tmp.replace(dst)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
