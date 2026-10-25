"""DOCX -> Google Docs import fidelity (master §29 "fidelity, ukuran, heading/table/link harus diuji").

The Google path publishes a DOCX snapshot and lets Drive convert it. Conversion is an implementation choice,
not a perfect-conversion claim, so every import is checked against what was actually sent:

* ``docx_structure`` reads the uploaded DOCX (python-docx): headings (level + text), tables (rows x columns),
  external hyperlinks, embedded pictures;
* ``gdoc_structure`` reads the same from the Docs API ``documents.get`` resource of the imported document;
* ``compare`` returns CONSISTENT only when all four match; any difference is MISMATCH with the reasons, and an
  unreadable document is UNKNOWN - never a pass.

The size limit is enforced before any request (``aeos.publisher.jobs.MAX_DOCX_BYTES``). Live behaviour on the
owner's Google account is exercised only by the owner-gated test ``tests/integration/test_google_docx_fidelity_live.py``.
"""

from __future__ import annotations

import io
import re
from typing import Any

HEADING_STYLE = re.compile(r"^Heading (\d)$")
GDOC_HEADING = re.compile(r"^HEADING_(\d)$")
CONSISTENT, MISMATCH, UNKNOWN = "CONSISTENT", "MISMATCH", "UNKNOWN"


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").replace("\u000b", " ")).strip()


def docx_structure(docx: bytes) -> dict[str, Any]:
    """Structure of the DOCX snapshot that is uploaded (the expected side)."""
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE

    doc = Document(io.BytesIO(docx))
    headings: list[list[Any]] = []
    for p in doc.paragraphs:
        m = HEADING_STYLE.match(p.style.name if p.style is not None else "")
        if m:
            headings.append([int(m.group(1)), _norm(p.text)])
    tables = [[len(t.rows), len(t.columns)] for t in doc.tables]
    # distinct link targets (python-docx shares one relationship per URL; Docs may split a link over text runs)
    links = sorted({r.target_ref for r in doc.part.rels.values()
                    if r.reltype == RELATIONSHIP_TYPE.HYPERLINK and r.is_external})
    return {"headings": headings, "tables": tables, "links": links, "images": len(doc.inline_shapes)}


def gdoc_structure(document: dict[str, Any]) -> dict[str, Any]:
    """Structure of a Docs API document resource (the observed side)."""
    headings: list[list[Any]] = []
    tables: list[list[int]] = []
    links: list[str] = []

    def walk(content: list[dict[str, Any]], *, top: bool) -> None:
        for el in content or []:
            if "paragraph" in el:
                para = el["paragraph"]
                text = "".join(pe.get("textRun", {}).get("content", "") for pe in para.get("elements", []))
                for pe in para.get("elements", []):
                    url = ((pe.get("textRun") or {}).get("textStyle") or {}).get("link", {}).get("url")
                    if url:
                        links.append(str(url))
                m = GDOC_HEADING.match(str((para.get("paragraphStyle") or {}).get("namedStyleType") or ""))
                if m and top:
                    headings.append([int(m.group(1)), _norm(text)])
            elif "table" in el:
                t = el["table"]
                tables.append([int(t.get("rows") or len(t.get("tableRows") or [])), int(t.get("columns") or 0)])
                for row in t.get("tableRows", []):
                    for cell in row.get("tableCells", []):
                        walk(cell.get("content", []), top=False)

    walk((document.get("body") or {}).get("content", []), top=True)
    return {"headings": headings, "tables": tables, "links": sorted(set(links)),
            "images": len(document.get("inlineObjects") or {})}


def compare(expected: dict[str, Any], observed: dict[str, Any] | None) -> dict[str, Any]:
    """CONSISTENT | MISMATCH | UNKNOWN with the differing aspects (headings, tables, links, images)."""
    counts = {k: len(expected.get(k) or []) if k != "images" else int(expected.get("images") or 0)
              for k in ("headings", "tables", "links", "images")}
    if observed is None:
        return {"status": UNKNOWN, "mismatches": ["imported document could not be read (documents.get)"],
                "expected": counts, "observed": None}
    obs_counts = {k: len(observed.get(k) or []) if k != "images" else int(observed.get("images") or 0)
                  for k in ("headings", "tables", "links", "images")}
    mismatches: list[str] = []
    if expected.get("headings") != observed.get("headings"):
        exp, got = expected.get("headings") or [], observed.get("headings") or []
        first = next((i for i, (a, b) in enumerate(zip(exp, got, strict=False)) if a != b), min(len(exp), len(got)))
        mismatches.append(f"headings: {len(exp)} expected, {len(got)} imported (first difference at #{first + 1}: "
                          f"{exp[first] if first < len(exp) else '-'} vs {got[first] if first < len(got) else '-'})")
    if expected.get("tables") != observed.get("tables"):
        mismatches.append(f"tables (rows x cols): {expected.get('tables')} expected, {observed.get('tables')} imported")
    if expected.get("links") != observed.get("links"):
        missing = sorted(set(expected.get("links") or []) - set(observed.get("links") or []))
        mismatches.append(f"links: {counts['links']} expected, {obs_counts['links']} imported"
                          + (f" (missing {', '.join(missing[:5])})" if missing else ""))
    if counts["images"] != obs_counts["images"]:
        mismatches.append(f"images: {counts['images']} embedded, {obs_counts['images']} imported")
    return {"status": MISMATCH if mismatches else CONSISTENT, "mismatches": mismatches, "expected": counts,
            "observed": obs_counts}
