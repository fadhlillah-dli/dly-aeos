"""DOCX snapshot of a rendered Product 1 document (master §29, S20, S21).

The Google Drive import path uploads this DOCX with conversion to Google Docs, which embeds rendered
images *inside the file* — no public/anyone-with-link image URLs are ever created (S21). Confluence
uses page storage format + attachments instead (see ``aeos.publisher.providers.atlassian``).

Supported subset of the canonical Markdown produced by ``aeos.docs_engine.renderer``:
headings (#..######), paragraphs, ``**bold**``, ``*italic*``/``_italic_``, ``code`` spans,
``[text](https://...)`` hyperlinks, bullet/numbered lists (one level), pipe tables, blockquotes,
fenced code blocks (monospace), images ``![alt](visual:<key>)`` resolved to rendered PNG files, and
section anchors (``<!-- aeos:section=... anchor=... -->``) as Word bookmarks on the heading.

Fidelity limits (documented, not hidden — §29 "bukan klaim konversi sempurna", §51 P2):
"""

from __future__ import annotations

import hashlib
import io
import re
from collections.abc import Callable
from pathlib import Path

from docx import Document as new_document
from docx.document import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

FIDELITY_LIMITS: tuple[str, ...] = (
    "Nested lists are flattened to one level.",
    "Table cells are plain text with inline formatting; no merged cells, widths or colours.",
    "Mermaid diagrams appear as images only when mmdc rendered them; otherwise a 'visual pending' note plus the "
    "source as monospace text (never a fake image).",
    "Only http(s)/mailto links become hyperlinks; other targets stay plain text.",
    "Section anchors become Word bookmarks; whether Google Docs keeps them after import is NOT-PROVEN.",
    "Google Docs conversion fidelity (heading/table/link/image) must be checked per tenant (master §51 P2).",
    "Control characters that XML 1.0 forbids are dropped from the snapshot (the Markdown source keeps them).",
)
__doc__ = (__doc__ or "") + "\n".join(f"* {x}" for x in FIDELITY_LIMITS) + "\n"

_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_IMAGE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
_ANCHOR = re.compile(r"<!--\s*aeos:section=\S+\s+anchor=(\S+)")
_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*\s][^*]*\*|_[^_\s][^_]*_)")
_SAFE_URL = re.compile(r"^(https?://|mailto:)", re.IGNORECASE)
# Characters XML 1.0 forbids (python-docx refuses them). Plan content is DATA and may contain them.
_XML_INVALID = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")

ImageResolver = Callable[[str], Path | None]


def _bookmark_name(anchor: str) -> str:
    # Word bookmark names: start with a letter, letters/digits/underscore, <= 40 chars.
    return "aeos_" + hashlib.sha256(anchor.encode("utf-8")).hexdigest()[:30]


def _add_bookmark(paragraph: Paragraph, anchor: str, bm_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bm_id))
    start.set(qn("w:name"), _bookmark_name(anchor))
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bm_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _add_hyperlink(paragraph: Paragraph, text: str, url: str) -> None:
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2F5BD3")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _add_inline(paragraph: Paragraph, text: str) -> None:
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**"):
            paragraph.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`"):
            r = paragraph.add_run(tok[1:-1])
            r.font.name = "Consolas"
        elif tok.startswith("["):
            label, _, url = tok[1:-1].partition("](")
            if _SAFE_URL.match(url.strip()):
                _add_hyperlink(paragraph, label, url.strip())
            else:
                paragraph.add_run(label)
        else:
            paragraph.add_run(tok[1:-1]).italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _style(doc: DocxDocument, name: str, fallback: str = "Normal") -> str:
    try:
        doc.styles[name]
        return name
    except KeyError:
        return fallback


def _split_row(line: str) -> list[str]:
    inner = line.strip().strip("|")
    cells = re.split(r"(?<!\\)\|", inner)
    return [c.strip().replace("\\|", "|") for c in cells]


def markdown_to_docx(markdown: str, *, title: str = "", logical_id: str = "",
                     image_resolver: ImageResolver | None = None) -> bytes:
    """Convert canonical renderer Markdown into DOCX bytes."""
    doc = new_document()
    if title:
        doc.core_properties.title = _XML_INVALID.sub("", title)[:250]
    if logical_id:
        doc.core_properties.keywords = f"aeos:logical_id={logical_id}"
    lines = _XML_INVALID.sub("", markdown.replace("\r\n", "\n")).split("\n")
    i = 0
    last_heading: Paragraph | None = None
    bm_id = 0
    bullet_style = _style(doc, "List Bullet")
    number_style = _style(doc, "List Number")
    quote_style = _style(doc, "Quote")
    table_style = _style(doc, "Table Grid", fallback="")
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("<!--"):
            m = _ANCHOR.search(stripped)
            if m and last_heading is not None:
                bm_id += 1
                _add_bookmark(last_heading, m.group(1), bm_id)
                last_heading = None
            i += 1
            continue
        if stripped.startswith("```"):
            i += 1
            code: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1  # closing fence
            for c in code or [""]:
                p = doc.add_paragraph(style=_style(doc, "No Spacing"))
                r = p.add_run(c)
                r.font.name = "Consolas"
                r.font.size = Pt(9)
            continue
        hm = _HEADING.match(stripped)
        if hm:
            level = min(len(hm.group(1)), 6)
            last_heading = doc.add_heading(level=level)
            _add_inline(last_heading, hm.group(2).strip())
            i += 1
            continue
        if stripped.startswith("|"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = _split_row(lines[i])
                if not all(re.fullmatch(r":?-{3,}:?", c or "---") for c in row):
                    rows.append(row)
                i += 1
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                if table_style:
                    table.style = doc.styles[table_style]
                for ri, row in enumerate(rows):
                    for ci in range(cols):
                        cell = table.cell(ri, ci)
                        para = cell.paragraphs[0]
                        _add_inline(para, row[ci] if ci < len(row) else "")
                        if ri == 0:
                            for run in para.runs:
                                run.bold = True
            continue
        im = _IMAGE.match(stripped)
        if im:
            alt, target = im.group(1), im.group(2).strip()
            path: Path | None = None
            if image_resolver is not None:
                key = target[len("visual:"):] if target.startswith("visual:") else target
                path = image_resolver(key)
            if path is not None and path.is_file():
                doc.add_picture(str(path), width=Inches(6))
                cap = doc.add_paragraph()
                cap.add_run(alt).italic = True
            else:
                p = doc.add_paragraph()
                r = p.add_run(f"[visual pending: {alt}]")
                r.italic = True
                r.font.color.rgb = RGBColor(0x8A, 0x5A, 0x00)
            i += 1
            continue
        if stripped.startswith(">"):
            quote: list[str] = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            text = " ".join(q for q in quote if q)
            if text:
                _add_inline(doc.add_paragraph(style=quote_style), text)
            continue
        bm = _BULLET.match(line)
        if bm:
            _add_inline(doc.add_paragraph(style=bullet_style), bm.group(1).strip())
            i += 1
            continue
        nm = _NUMBERED.match(line)
        if nm:
            _add_inline(doc.add_paragraph(style=number_style), nm.group(1).strip())
            i += 1
            continue
        para_lines = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not any(
            lines[i].strip().startswith(x) for x in ("#", "|", ">", "```", "<!--", "![")
        ) and not _BULLET.match(lines[i]) and not _NUMBERED.match(lines[i]):
            para_lines.append(lines[i].strip())
            i += 1
        _add_inline(doc.add_paragraph(), " ".join(para_lines))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
