"""DOCX snapshots (§29 Drive import path, S20/S21) and visuals (mermaid-cli optional, never fake images)."""

from __future__ import annotations

import base64
import copy
import io
import subprocess
import zipfile

import docx as python_docx
import pytest

from aeos.docs_engine import renderer as r
from aeos.docs_engine import visuals as vis
from aeos.docs_engine.docx import FIDELITY_LIMITS, markdown_to_docx
from aeos.docs_engine.plan_content import DiagramView
from tests.unit.test_publisher_support import SAMPLE_CONTENT

# 1x1 transparent PNG
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")

MD = """# Judul Dokumen

<!-- aeos:logical_id=x renderer=1 -->

| Metadata | Nilai |
|---|---|
| Logical doc ID | `x` |

## Bagian Satu

<!-- aeos:section=one anchor=x--one content_type=AI_MAINTAINED -->
Paragraf dengan **tebal**, *miring*, `kode` dan [tautan](https://example.org/a) serta [bukan](javascript:alert(1)).

- butir satu
- butir dua

1. langkah satu

> catatan penting

```mermaid
flowchart TD
  A --> B
```

![Diagram A](visual:diag-a)

![Diagram B](visual:missing)
"""


def _doc(data: bytes):
    return python_docx.Document(io.BytesIO(data))


def test_docx_contains_headings_tables_links_and_images(tmp_path):
    img = tmp_path / "a.png"
    img.write_bytes(PNG)
    data = markdown_to_docx(MD, title="Judul", logical_id="x", image_resolver=lambda k: img if k == "diag-a" else None)
    d = _doc(data)
    styles = [(p.style.name, p.text) for p in d.paragraphs]
    assert ("Heading 1", "Judul Dokumen") in styles and ("Heading 2", "Bagian Satu") in styles
    assert any(s == "List Bullet" and t == "butir satu" for s, t in styles)
    assert any(s == "List Number" for s, _t in styles)
    assert any("[visual pending: Diagram B]" in t for _s, t in styles)
    assert d.tables and d.tables[0].cell(1, 1).text == "x"
    assert d.core_properties.keywords == "aeos:logical_id=x"
    z = zipfile.ZipFile(io.BytesIO(data))
    rels = z.read("word/_rels/document.xml.rels").decode()
    assert "https://example.org/a" in rels and "javascript:" not in rels  # only safe links
    assert any(n.startswith("word/media/") for n in z.namelist())  # image embedded in the file (no public URL)
    body = z.read("word/document.xml").decode()
    assert "w:bookmarkStart" in body and "flowchart TD" in body


def test_docx_of_a_rendered_document_loads():
    res = r.render_document_set(
        project=r.ProjectIdentity(id="p", key="k", name="Contoh", domain="PERSONAL"),
        version=r.VersionIdentity(id="v", plan_id="pl", version=1, manifest_hash="sha256:" + "c" * 64, approved=True),
        content=copy.deepcopy(SAMPLE_CONTENT))
    tp = res.by_id()["06-testing/test-plan-tasks"]
    d = _doc(markdown_to_docx(tp.body_markdown or "", title=tp.title))
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert "Test Strategy" in heads and "Sign-off" in heads


def test_fidelity_limits_are_documented():
    assert any("NOT-PROVEN" in x for x in FIDELITY_LIMITS)
    import aeos.docs_engine.docx as mod

    assert "Fidelity limits" in (mod.__doc__ or "")


@pytest.fixture
def no_mmdc(monkeypatch):
    monkeypatch.delenv("AEOS_MMDC_BINARY", raising=False)
    monkeypatch.setattr(vis.shutil, "which", lambda _name: None)


DIAG = DiagramView(key="ctx", title="Context", kind="context", source="flowchart TD\n  A-->B")


def test_missing_mmdc_keeps_source_and_marks_visual_pending(tmp_path, no_mmdc):
    info = vis.detect_mmdc()
    assert not info.available and info.status == "NOT_INSTALLED" and "mermaid-cli" in info.reason
    [art] = vis.render_visuals([DIAG], tmp_path)
    assert art.status == vis.STATUS_PENDING and art.png_path is None and art.svg_path is None
    assert not list(tmp_path.glob("*.png")) and list(tmp_path.glob("*.mmd"))  # source kept, no fake image
    assert art.evidence_role == "ILLUSTRATIVE_NOT_RUNTIME_PROOF" and art.design_status == "DESIGN"


def test_mmdc_render_success_and_cache(tmp_path, monkeypatch):
    fake = tmp_path / "mmdc.cmd"
    fake.write_text("x")
    calls = []

    def run(args, **kw):
        calls.append(args)
        assert isinstance(args, list) and kw.get("timeout") and "shell" not in kw  # bounded, list args, no shell
        if "--version" in args:
            return subprocess.CompletedProcess(args, 0, "11.4.0\n", "")
        out = args[args.index("-o") + 1]
        with open(out, "wb") as fh:
            fh.write(PNG)
        return subprocess.CompletedProcess(args, 0, "", "")

    monkeypatch.setattr(vis, "run_bounded", run)
    monkeypatch.setenv("AEOS_MMDC_BINARY", str(fake))
    [art] = vis.render_visuals([DIAG], tmp_path / "v")
    assert art.status == vis.STATUS_RENDERED and art.png_path and art.svg_path and "11.4.0" in (art.renderer or "")
    n = len(calls)
    [again] = vis.render_visuals([DIAG], tmp_path / "v")
    assert again.status == vis.STATUS_RENDERED and len(calls) == n + 1  # only --version, render cached by hash


def test_mmdc_timeout_is_failed_not_fake(tmp_path, monkeypatch):
    fake = tmp_path / "mmdc"
    fake.write_text("x")

    def run(args, **kw):
        if "--version" in args:
            return subprocess.CompletedProcess(args, 0, "11\n", "")
        raise subprocess.TimeoutExpired(args, kw["timeout"])

    monkeypatch.setattr(vis, "run_bounded", run)
    monkeypatch.setenv("AEOS_MMDC_BINARY", str(fake))
    [art] = vis.render_visuals([DIAG], tmp_path / "v", )
    assert art.status == vis.STATUS_FAILED and "timeout" in art.reason and art.png_path is None


def test_mmdc_timeout_kills_the_whole_process_tree(tmp_path):
    """mmdc is an npm shim on Windows (mmdc.cmd -> node -> Chromium): a timeout must not leave children."""
    import sys
    import time

    import psutil

    pid_file = tmp_path / "child.pid"
    script = ("import subprocess, sys, time; "
              "c = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)']); "
              f"open(r'{pid_file}', 'w').write(str(c.pid)); time.sleep(60)")
    started = time.monotonic()
    with pytest.raises(subprocess.TimeoutExpired):
        vis.run_bounded([sys.executable, "-c", script], timeout=6)
    # timing-cap: 30 s vs a measured 6.0 s (the 6 s timeout plus 0.01 s); a missing kill waits the 60 s sleep
    assert time.monotonic() - started < 30
    child = int(pid_file.read_text())
    deadline = time.monotonic() + 10
    while psutil.pid_exists(child) and time.monotonic() < deadline:
        try:
            if psutil.Process(child).status() == psutil.STATUS_ZOMBIE:
                break
        except psutil.Error:
            break
        time.sleep(0.1)
    alive = psutil.pid_exists(child)
    if alive:
        try:
            alive = psutil.Process(child).status() != psutil.STATUS_ZOMBIE
        except psutil.Error:
            alive = False
    assert not alive, "grandchild survived the timeout"


def test_renderer_embeds_rendered_or_pending_visuals(tmp_path, no_mmdc):
    pending = vis.render_visuals([DIAG], tmp_path)
    res = r.render_document_set(
        project=r.ProjectIdentity(id="p", key="k", name="Contoh", domain="PERSONAL"),
        version=r.VersionIdentity(id="v", plan_id="pl", version=1, manifest_hash="sha256:" + "d" * 64, approved=True),
        content=copy.deepcopy(SAMPLE_CONTENT), visuals=pending)
    arch = res.by_id()["03-architecture-overview"].body_markdown or ""
    assert "Visual pending" in arch and "```mermaid" in arch and "bukan bukti runtime" in arch
    assert "![" not in arch


def test_screenshot_manifest_is_not_proven():
    m = vis.screenshot_manifest(project_id="p", commit="abc123", environment="local", routes=["/"])
    assert m["status"] == "NOT_PROVEN" and m["items"] == [] and m["commit"] == "abc123"
    assert any("secret" in x.lower() for x in m["requirements"])
