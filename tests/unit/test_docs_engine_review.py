"""Docs engine: §27 metadata (last reviewed), §28 feature-approval threat/assessment, safe logical ids,
local preview tree and CLI audit (§27-§29, §41).

§50 row: "VIEWER memanggil write endpoint langsung → Ditolak backend" (document review).
"""

from __future__ import annotations

import copy
import io

import docx as python_docx
import pytest
from sqlalchemy import select
from typer.testing import CliRunner

from aeos.core.enums import Domain
from aeos.core.paths import safe_segment
from aeos.docs_engine import product1 as p1
from aeos.docs_engine import renderer as r
from aeos.docs_engine import service
from aeos.docs_engine.docx import markdown_to_docx
from aeos.models.audit import AuditEvent
from aeos.models.planning import Plan
from aeos.models.publishing import DocumentRecord
from tests.conftest import VIEWER_PW, create_admin, create_viewer
from tests.unit.test_publisher_support import (
    SAMPLE_CONTENT,
    add_version,
    approve,
    install_fakes,
    make_plan,
    make_project,
)

PROJECT = r.ProjectIdentity(id="p-1", key="contoh", name="Contoh App", domain="PERSONAL", owner_name="Owner")
VERSION = r.VersionIdentity(id="v-1", plan_id="plan-1", version=1, manifest_hash="sha256:" + "a" * 64, approved=True)


def _render(content):
    return r.render_document_set(project=PROJECT, version=VERSION, content=content)


# --------------------------------------------------------------------------------------------- review


@pytest.fixture
def world(settings, db, secrets_memory, monkeypatch):
    install_fakes(monkeypatch)
    project = make_project(db)
    plan, version = make_plan(db, project)
    return {"db": db, "project": project, "plan": plan, "version": version}


def _render_api(client, world, version_id=None):
    r_ = client.post(f"/api/projects/{world['project'].id}/documents/render",
                     json={"plan_version_id": version_id or world["version"].id, "render_visuals": False})
    assert r_.status_code == 200, r_.text


def _item(client, world, logical_id):
    docs = client.get(f"/api/projects/{world['project'].id}/documents").json()
    return next(i for i in docs["items"] if i["logical_doc_id"] == logical_id)


def test_review_is_bound_to_the_exact_content_hash_and_goes_stale(personal_client, world):
    _render_api(personal_client, world)
    item = _item(personal_client, world, "01-product-overview")
    assert item["review"]["status"] == "NOT_REVIEWED"
    stale = personal_client.post(f"/api/documents/{item['id']}/review", json={"content_hash": "sha256:old"})
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "STALE_CONTENT"
    ok = personal_client.post(f"/api/documents/{item['id']}/review",
                              json={"content_hash": item["content_hash"], "note": "dibaca"})
    assert ok.status_code == 200 and ok.json()["status"] == "REVIEWED"
    item = _item(personal_client, world, "01-product-overview")
    assert item["review"]["status"] == "REVIEWED" and item["last_reviewed_at"]
    db = world["db"]
    db.expire_all()
    ev = db.scalar(select(AuditEvent).where(AuditEvent.action == "document.review"))
    assert ev is not None and ev.target_id == item["id"]
    # a newer approved version changes the document → the old review is visibly out of date
    plan = db.get(Plan, world["plan"].id)
    content = copy.deepcopy(SAMPLE_CONTENT)
    content["summary"] = "Ringkasan baru setelah review"
    v2 = add_version(db, plan, content)
    approve(db, plan, v2)
    db.commit()
    _render_api(personal_client, world, v2.id)
    assert _item(personal_client, world, "01-product-overview")["review"]["status"] == "REVIEW_STALE"
    folder = _item(personal_client, world, "02-requirements")
    nf = personal_client.post(f"/api/documents/{folder['id']}/review", json={"content_hash": "x"})
    assert nf.status_code == 409 and nf.json()["error"]["code"] == "NOT_A_DOCUMENT"


def test_viewer_cannot_record_a_review(client, db, world):
    create_admin(db)
    create_viewer(db)
    rec_id = "doc-x"
    client.login("viewer", VIEWER_PW)
    client.select_domain(Domain.PERSONAL)
    res = client.post(f"/api/documents/{rec_id}/review", json={"content_hash": "sha256:x"})
    assert res.status_code == 403 and res.json()["error"]["code"] == "READ_ONLY_ROLE"


def test_review_of_another_domains_document_is_404(personal_client, db, world, monkeypatch):
    office = make_project(db, domain="OFFICE", name="Kantor")
    _plan, ov = make_plan(db, office)
    from aeos.core.config import get_settings

    service.render_version(db, get_settings(), ov, persist=True, render_visuals=False)
    db.commit()
    rec = db.scalar(select(DocumentRecord).where(DocumentRecord.project_id == office.id,
                                                 DocumentRecord.logical_doc_id == "00-product-home"))
    res = personal_client.post(f"/api/documents/{rec.id}/review", json={"content_hash": rec.content_hash})
    assert res.status_code == 404


def test_document_body_never_claims_a_review_state():
    home = _render(copy.deepcopy(SAMPLE_CONTENT)).by_id()["00-product-home"].body_markdown
    assert "belum direview" not in home and "Last reviewed" in home


# ------------------------------------------------------------------------------- feature approval (§28)


def _with_security(content):
    content = copy.deepcopy(content)
    f = content["features"][1]  # auth
    f["security_assessment"] = {"scope": "Login lokal", "requirements": ["Argon2id", "Rate limit"],
                                "risks": ["Credential stuffing"], "mitigations": ["Lockout"]}
    f["threat_model"] = {"assets": ["Password hash"], "actors": ["Penyerang online"], "entry_points": ["/login"]}
    f["threats"] = ["Brute force"]
    return content


def test_feature_approval_renders_security_assessment_and_threat_model_per_feature():
    res = _render(_with_security(SAMPLE_CONTENT))
    ids = res.by_id()
    sa = ids["04-security-assessments/security-assessment-auth"]
    tm = ids["04-threat-models/threat-model-auth"]
    assert sa.parent_logical_id == "04-security-assessments" and tm.parent_logical_id == "04-threat-models"
    assert sa.timing == tm.timing == "FEATURE_APPROVAL" and sa.template_key == "security-assessment"
    secs = {s.key: s.status for s in sa.sections}
    assert secs["findings"] == r.SEC_LIVE and secs["residual_risk"] == secs["approval"] == r.SEC_HUMAN
    assert secs["requirements"] == r.SEC_FILLED and "Argon2id" in sa.body_markdown
    assert "Temuan hanya dari review/scan nyata" in sa.body_markdown and "PASS" not in sa.body_markdown
    assert "Brute force" in tm.body_markdown and "Password hash" in tm.body_markdown
    assert "04-security-assessments/security-assessment-tasks" not in ids  # nothing requested → nothing invented
    assert "04-threat-models/threat-model-tasks" not in ids
    # the initial (project-level) threat model still exists next to the feature one
    assert "04-threat-models/threat-model-initial" in ids
    # the flag alone creates the document with pending sections — never invented content
    content = copy.deepcopy(SAMPLE_CONTENT)
    content["features"][0]["needs_security_assessment"] = True
    flagged = _render(content).by_id()["04-security-assessments/security-assessment-tasks"]
    assert "Requirements" in flagged.pending_sections and "Risks" in flagged.pending_sections


def test_threat_models_folder_has_items_when_only_a_feature_threat_model_exists():
    content = _with_security(SAMPLE_CONTENT)
    content["security"] = {"authn": "Password"}  # no project-level threats
    ids = _render(content).by_id()
    assert "04-threat-models/threat-model-initial" not in ids
    assert ids["04-threat-models"].status == "APPROVED" and "04-threat-models/threat-model-auth" in ids


# ------------------------------------------------------------------------------------------ safe ids


def test_rfc_and_tdd_keys_from_plan_data_are_slugged_and_never_collide():
    content = copy.deepcopy(SAMPLE_CONTENT)
    content["rfcs"] = [{"key": "../../etc/passwd", "title": "A"}, {"key": "../../etc/passwd", "title": "B"},
                       {"key": "!!!", "title": "C"}, {"key": "auth", "title": "Plan-level auth RFC"}]
    res = _render(content)
    rfcs = [d for d in res.documents if d.doc_type == "RFC"]
    lids = [d.logical_doc_id for d in rfcs]
    assert lids == ["03-rfcs/rfc-auth", "03-rfcs/rfc-etc-passwd", "03-rfcs/rfc-etc-passwd-2", "03-rfcs/rfc-rfc-3",
                    "03-rfcs/rfc-auth-2"]
    for lid in lids:
        for seg in lid.split("/"):
            safe_segment(seg)  # raises on traversal/separators
    # the feature-bound RFC keeps the id the PRD links to
    prd = res.by_id()["02-requirements/prd-auth"]
    assert "03-rfcs/rfc-auth" in prd.links["related"]


# ------------------------------------------------------------------------------------- local preview


def test_local_preview_mirrors_the_nested_product1_tree(settings, db, monkeypatch, tmp_path):
    install_fakes(monkeypatch)
    project = make_project(db)
    _plan, v1 = make_plan(db, project)
    res, _d, _s = service.render_version(db, settings, v1, persist=False, render_visuals=False)
    out = tmp_path / "out"
    info = service.write_preview(db, res, out)
    assert (out / "04-security-privacy" / "04-threat-models" / "threat-model-initial.md").is_file()
    assert (out / "03-architecture-design" / "03-architecture-overview.md").is_file()
    assert (out / "templates" / "prd.md").is_file()
    assert "04-security-privacy/04-threat-models/threat-model-initial.md" in info["files"]
    assert all("\\" not in f for f in info["files"])  # POSIX-style relative paths on every OS
    index = (out / "INDEX.md").read_text(encoding="utf-8")
    assert "\n" + "  " * 3 + "- [doc] Threat Model — Initial" in index  # depth follows the folder nesting
    assert "\n" + "  " * 1 + "- [doc] 00. Product Home" in index


# ------------------------------------------------------------------------------------------------ CLI


def test_cli_render_of_the_approved_version_is_audited(settings, db, monkeypatch, tmp_path):
    install_fakes(monkeypatch)
    project = make_project(db)
    plan, v1 = make_plan(db, project)
    from aeos.cli import app

    res = CliRunner().invoke(app, ["docs", "render", v1.id, "--out", str(tmp_path / "a"), "--no-visuals"])
    assert res.exit_code == 0, res.output
    db.expire_all()
    ev = db.scalar(select(AuditEvent).where(AuditEvent.action == "documents.render"))
    assert ev is not None and ev.actor_type == "cli" and ev.project_id == project.id
    # an unapproved version is a local preview only: no DocumentRecord change, no audit entry
    v2 = add_version(db, db.get(Plan, plan.id), {"summary": "draft"})
    db.commit()
    res2 = CliRunner().invoke(app, ["docs", "render", v2.id, "--out", str(tmp_path / "b"), "--no-visuals"])
    assert res2.exit_code == 0 and "local preview only" in res2.output
    db.expire_all()
    assert len(list(db.scalars(select(AuditEvent).where(AuditEvent.action == "documents.render")))) == 1


# ----------------------------------------------------------------------------------------------- DOCX


def test_docx_snapshot_survives_xml_invalid_characters_in_plan_data():
    data = markdown_to_docx("# Judul\x01\n\nTeks dengan \x0b kontrol\x00 karakter\n\n| A\x02 | B |\n|---|---|\n| 1 | 2 |\n",
                            title="Judul\x07")
    d = python_docx.Document(io.BytesIO(data))
    texts = [p.text for p in d.paragraphs]
    assert "Judul" in texts and any("kontrol karakter" in t for t in texts)
    assert d.tables and d.tables[0].cell(0, 0).text == "A"


def test_every_structure_section_title_is_unique_per_document_type():
    for spec in p1.DOC_TYPES.values():
        keys = [s.key for s in spec.sections]
        assert len(keys) == len(set(keys)), spec.doc_type
